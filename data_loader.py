from __future__ import annotations
import io
import pandas as pd
import numpy as np
from openpyxl import load_workbook


def make_template_csv() -> bytes:
    df = pd.DataFrame({
        "x": [1.0, 1.0, 1.5, 1.5, 2.0, 2.0],
        "y": [10.2, 12.1, 18.4, 19.0, 24.2, 25.1],
        "group": ["A", "A", "B", "B", "C", "C"],
        "condition": ["Condition A", "Condition A", "Condition B", "Condition B", "Condition C", "Condition C"],
    })
    return df.to_csv(index=False).encode("utf-8-sig")


def _read_csv(raw: bytes) -> pd.DataFrame:
    for enc in ("utf-8-sig", "utf-8", "cp932", "shift_jis"):
        try:
            return pd.read_csv(io.BytesIO(raw), encoding=enc)
        except Exception:
            pass
    raise ValueError("CSV encoding could not be detected.")


def load_uploaded_file(uploaded) -> pd.DataFrame:
    name = uploaded.name.lower()
    raw = uploaded.getvalue()
    if name.endswith(".csv"):
        df = _read_csv(raw)
    elif name.endswith((".xlsx", ".xls")):
        xls = pd.ExcelFile(io.BytesIO(raw))
        df = pd.read_excel(io.BytesIO(raw), sheet_name=xls.sheet_names[0])
    else:
        raise ValueError("Unsupported file type.")
    df = df.dropna(how="all").dropna(axis=1, how="all")
    df.columns = [str(c).strip() for c in df.columns]
    return df


def _parse_block(ws, before_col, after_col, start_row, end_row, condition, grain_size):
    records = []
    indent = ""
    crack_no = 0
    bcol = ws[f"{before_col}{start_row}"].column
    for r in range(start_row, end_row + 1):
        for c in range(max(1, bcol - 3), min(ws.max_column, bcol + 4) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith("No."):
                indent = v
                crack_no = 0
                break
        before = ws[f"{before_col}{r}"].value
        after = ws[f"{after_col}{r}"].value
        if isinstance(before, (int, float)) and isinstance(after, (int, float)) and before != 0:
            crack_no += 1
            records.append({
                "grain_size_um": float(grain_size),
                "healing_pct": 100.0 * (1.0 - float(after) / float(before)),
                "condition": condition,
                "indentation": indent,
                "crack": f"L{crack_no}",
                "before_um": float(before),
                "after_um": float(after),
            })
    return records


def load_crack_healing_preset(uploaded):
    raw = uploaded.getvalue()
    wb = load_workbook(io.BytesIO(raw), data_only=False)
    sheet = next((s for s in wb.sheetnames if "950" in s), wb.sheetnames[0])
    ws = wb[sheet]
    specs = [
        ("D", "H", 33, 56, "1550°C, 1 h", 1.296037333),
        ("O", "S", 33, 56, "1550°C, 12 h", 1.531909333),
        ("Y", "AC", 33, 64, "1650°C, 1 h", 1.632668150),
        ("AJ", "AN", 33, 64, "1650°C, 12 h", 2.026158000),
    ]
    records = []
    for args in specs:
        records.extend(_parse_block(ws, *args))
    df = pd.DataFrame(records).replace([np.inf, -np.inf], np.nan).dropna(subset=["healing_pct", "grain_size_um"])
    if df.empty:
        raise ValueError("No valid crack-healing data were found.")
    return df, f"Crack Healing preset loaded: {len(df)} crack measurements."
